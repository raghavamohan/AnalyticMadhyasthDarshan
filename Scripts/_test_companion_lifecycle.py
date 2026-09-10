"""Exercise companion registration, fresh outputs and lifecycle manifest changes."""
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from pptx import Presentation

import _companion_lifecycle as lifecycle
import _presentation_pipeline as pipeline
import _verify_companion_outputs as companions


class CompanionLifecycleTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name).resolve()
        self.source = self.root / 'Studies/A/Presenters-Companion-Deck.md'
        self.source.parent.mkdir(parents=True)
        self.source.write_text('# Slide 1\n\n## Delivering the slide\n\nThe delivery script.\n', encoding='utf-8')
        self.deck = self.source.parent / 'Deck.pptx'
        pptx = Presentation()
        pptx.slides.add_slide(pptx.slide_layouts[6])
        pptx.save(self.deck)
        self.row = {'markdown': self.source.relative_to(self.root).as_posix(), 'deck': 'a'}
        self.write_companions([self.row])
        self.deck_row = {'id': 'a', 'source': 'Studies/A/Deck.pptx',
                         'slidesPdf': 'Studies/A/Deck.pdf', 'notesPdf': 'Studies/A/Deck-notes.pdf'}
        self.presentation_path = self.root / 'Scripts/presentation-pipeline.json'
        self.presentation_path.write_text(json.dumps({'productionProfile': 'keep', 'decks': [self.deck_row]}), encoding='utf-8')
        spec = SimpleNamespace(id='a', source=self.deck)
        manifest = SimpleNamespace(decks=[spec], deck=lambda _id: spec)
        base = patch.object(companions, 'BASE', self.root)
        loader = patch.object(pipeline, 'load_manifest', return_value=manifest)
        base.start()
        loader.start()
        self.addCleanup(base.stop)
        self.addCleanup(loader.stop)

    def write_companions(self, rows):
        path = self.root / 'Scripts/companion-pipeline.json'
        path.parent.mkdir(exist_ok=True)
        path.write_text(json.dumps({'schema': 1, 'companions': rows}), encoding='utf-8')
        return path

    def test_new_companion_requires_registration(self):
        self.write_companions([])
        with self.assertRaisesRegex(ValueError, 'Unregistered presenter'):
            companions.load_companions()

    def test_rejects_unsafe_missing_duplicate_and_cross_study_ownership(self):
        for rows in ([{'markdown': '../outside.md', 'deck': 'a'}],
                     [{'markdown': 'Studies/A/Presenters-Companion-Missing.md', 'deck': 'a'}],
                     [self.row, self.row], [dict(self.row, deck='unknown')], [None]):
            with self.subTest(rows=rows):
                self.write_companions(rows)
                with self.assertRaises(ValueError):
                    companions.load_companions()
        other = self.root / 'Studies/B/Presenters-Companion-Deck.md'
        other.parent.mkdir()
        other.write_bytes(self.source.read_bytes())
        self.write_companions([{'markdown': other.relative_to(self.root).as_posix(), 'deck': 'a'}])
        with self.assertRaisesRegex(ValueError, 'same study'):
            companions.load_companions()

    def test_preparation_and_freshness_use_the_same_markdown_chain(self):
        self.assertTrue(companions.verify())  # Missing generated notes and DOCX.
        companions.prepare({self.row['markdown']})
        self.assertEqual(companions.verify(), [])
        paths = [self.deck, self.source.with_suffix('.docx'), self.source.with_suffix('.notes.json')]
        before = [p.read_bytes() for p in paths]
        companions.prepare({self.row['markdown']})
        self.assertEqual([p.read_bytes() for p in paths], before)
        self.source.write_text(self.source.read_text(encoding='utf-8').replace('The delivery', 'An updated delivery'), encoding='utf-8')
        errors = '\n'.join(companions.verify())
        self.assertIn('notes.json differs', errors)
        self.assertIn('speaker notes differ', errors)
        self.assertIn('.docx differs', errors)
        companions.prepare({self.row['markdown']})
        self.assertEqual(companions.verify(), [])

    def test_rename_updates_both_manifests_and_preserves_other_study_and_ids(self):
        companion_path = self.write_companions([self.row, {'markdown': 'Studies/AB/Presenters-Companion-Deck.md', 'deck': 'ab'}])
        before = (self.presentation_path.read_bytes(), companion_path.read_bytes())
        self.assertEqual(lifecycle.rename_paths('A', 'New', root=self.root, dry_run=True), 4)
        self.assertEqual((self.presentation_path.read_bytes(), companion_path.read_bytes()), before)
        self.assertEqual(lifecycle.rename_paths('A', 'New', root=self.root), 4)
        deck = json.loads(self.presentation_path.read_bytes())['decks'][0]
        self.assertEqual(deck['source'], 'Studies/New/Deck.pptx')
        self.assertEqual(deck['slidesPdf'], 'Studies/New/Deck.pdf')
        self.assertEqual(deck['notesPdf'], 'Studies/New/Deck-notes.pdf')
        self.assertEqual(deck['id'], 'a')
        data = json.loads(companion_path.read_bytes())['companions']
        self.assertEqual(data[0]['markdown'], 'Studies/New/Presenters-Companion-Deck.md')
        self.assertEqual(data[1]['markdown'], 'Studies/AB/Presenters-Companion-Deck.md')
        self.assertEqual(lifecycle.rename_paths('A', 'New', root=self.root), 0)

    def test_retirement_handles_applied_metadata_only_and_is_idempotent(self):
        path = self.write_companions([dict(self.row, markdown='Applications/Gone/Presenters-Companion-Deck.md'), self.row])
        original = path.read_bytes()
        self.assertEqual(lifecycle.remove_companions('Gone', root=self.root, dry_run=True), 1)
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(lifecycle.remove_companions('Gone', root=self.root), 1)
        self.assertEqual(json.loads(path.read_bytes())['companions'], [self.row])
        self.assertEqual(lifecycle.remove_companions('Gone', root=self.root), 0)

    def test_prepared_manifest_cannot_redefine_binary_ownership(self):
        from _prepared_study import validate
        from _verification_identity import intent_hash
        pr = {'number': 1, 'state': 'open', 'draft': True, 'body': '',
              'head': {'sha': 'a' * 40, 'repo': {'full_name': 'owner/repo'}},
              'base': {'ref': 'master', 'sha': 'b' * 40}}
        payload = {'schema': 1, 'pr': 1, 'repository': 'owner/repo', 'head': 'a' * 40,
                   'base': 'b' * 40, 'intent': intent_hash(pr), 'files': {}}
        import base64
        with patch('_prepared_study.BASE', self.root):
            for rows in ([], [self.row], [dict(self.row, markdown='Studies/New/Presenters-Companion-Deck.md')]):
                raw = json.dumps({'schema': 1, 'companions': rows}).encode()
                payload['files'] = {'Scripts/companion-pipeline.json': base64.b64encode(raw).decode()}
                self.assertEqual(validate(payload, pr, 'owner/repo')['Scripts/companion-pipeline.json'], raw)
            for data in ({'schema': 2, 'companions': []},
                         {'schema': 1, 'companions': [dict(self.row, deck='new-owner')]},
                         {'schema': 1, 'companions': [dict(self.row, markdown='Scripts/Arbitrary.md')]},
                         {'schema': 1, 'companions': [dict(self.row, markdown='Studies/A/Different.md')]},
                         {'schema': 1, 'companions': [self.row, self.row]}):
                with self.subTest(data=data):
                    payload['files'] = {'Scripts/companion-pipeline.json': base64.b64encode(json.dumps(data).encode()).decode()}
                    with self.assertRaises(ValueError):
                        validate(payload, pr, 'owner/repo')


if __name__ == '__main__':
    unittest.main()
